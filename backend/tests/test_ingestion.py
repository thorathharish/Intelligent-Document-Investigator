"""Checkpoint 2: upload -> validate -> extract -> clean -> chunk -> embed -> persist."""
import io

import docx
import pymupdf
import pytest
from fastapi.testclient import TestClient

from app import db
from app.ingestion import pipeline
from app.ingestion.chunk import MAX_PARAGRAPH_CHARS, chunk_pages
from app.ingestion.extract import Page, Paragraph
from app.ingestion.validate import ValidationError, validate_file
from app.main import app


@pytest.fixture(scope="module")
def client():
    with TestClient(app) as c:
        yield c


def _pdf_bytes(pages: list[list[tuple[str, str]]], body_size: float = 10.5) -> bytes:
    doc = pymupdf.open()
    for items in pages:
        page = doc.new_page()
        y = 72.0
        for kind, text in items:
            rect = pymupdf.Rect(72, y, 523, 800)
            left = page.insert_textbox(
                rect, text, fontsize=13 if kind == "h" else body_size, fontname="hebo" if kind == "h" else "helv"
            )
            y += rect.height - left + 14
    data = doc.tobytes()
    doc.close()
    return data


PDF = _pdf_bytes(
    [
        [("h", "1. Overview"), ("p", "ALPHAPAGE This agreement describes the services supplied to the customer by the supplier.")],
        [("h", "4.2 Payment Terms"), ("p", "BETAPAGE The Customer shall pay each undisputed invoice within thirty (30) days of the invoice date.")],
    ]
)
TXT = b"HR MEMO JULY\n\nEffective 1 July 2025, remote work is not permitted for any employee of the company.\n"


def _docx_bytes() -> bytes:
    document = docx.Document()
    document.add_heading("Invoice Approval", level=1)
    document.add_paragraph("Invoices above USD 25,000 must be approved by the Finance Director before payment.")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Late fee"
    table.rows[0].cells[1].text = "1.5% per month on overdue vendor invoices"
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _png_bytes() -> bytes:
    doc = pymupdf.open(stream=_pdf_bytes([[("p", "Payment due within 45 days of the invoice date.")]], body_size=16), filetype="pdf")
    data = doc[0].get_pixmap(dpi=150).tobytes("png")
    doc.close()
    return data


def _new_investigation(client) -> str:
    return client.post("/api/investigations", json={"title": "Test"}).json()["id"]


def _upload(client, investigation_id: str, files: list[tuple[str, bytes]]):
    return client.post(
        f"/api/investigations/{investigation_id}/documents",
        files=[("files", (name, data, "application/octet-stream")) for name, data in files],
    )


def _documents(client, investigation_id: str) -> dict[str, dict]:
    return {d["filename"]: d for d in client.get(f"/api/investigations/{investigation_id}/documents").json()}


def _chunks(document_id: str):
    return db.query("SELECT * FROM chunks WHERE document_id = ? ORDER BY ordinal", (document_id,))


# --- validation ---------------------------------------------------------------


def test_valid_file_returns_ext_and_hash():
    ext, sha = validate_file("Contract.PDF", PDF)
    assert ext == "pdf" and len(sha) == 64


def test_unsupported_extension():
    with pytest.raises(ValidationError, match="Unsupported"):
        validate_file("sheet.xlsx", b"PK\x03\x04data")


def test_oversized_file():
    with pytest.raises(ValidationError, match="larger"):
        validate_file("notes.txt", b"x" * 100, max_bytes=10)


def test_oversized_upload_is_rejected_without_being_read():
    from app.api.documents import read_capped

    class Unreadable:
        def read(self, *args):
            raise AssertionError("an oversized file must not be read into memory")

    class Recording:
        def __init__(self, data):
            self.data, self.asked = data, None

        def read(self, n=-1):
            self.asked = n
            return self.data[:n]

    class Upload:
        def __init__(self, filename, size, file):
            self.filename, self.size, self.file = filename, size, file

    limit = 20 * 1024 * 1024
    name, data, size = read_capped(Upload("huge.pdf", 5 * 1024**3, Unreadable()), limit)
    assert (name, data, size) == ("huge.pdf", b"", 5 * 1024**3)
    with pytest.raises(ValidationError, match="larger"):
        validate_file(name, data, size=size)

    # size unknown: never more than limit + 1 bytes are read, which is enough to see it is too large
    source = Recording(b"x" * 64)
    name, data, size = read_capped(Upload("notes.txt", None, source), 16)
    assert source.asked == 17 and size == 17
    with pytest.raises(ValidationError, match="larger"):
        validate_file(name, data, max_bytes=16, size=size)

    # a normal file is read whole and validated as before
    name, data, size = read_capped(Upload("memo.txt", len(TXT), Recording(TXT)), limit)
    assert data == TXT and validate_file(name, data, size=size)[0] == "txt"


def test_oversized_file_in_an_upload_does_not_block_the_others(client, monkeypatch):
    from app.config import settings

    monkeypatch.setattr("app.api.documents.settings", settings)
    original = settings.max_file_mb
    object.__setattr__(settings, "max_file_mb", 1)
    try:
        inv = _new_investigation(client)
        body = _upload(client, inv, [("big.txt", b"word " * 300_000), ("memo.txt", TXT)]).json()
    finally:
        object.__setattr__(settings, "max_file_mb", original)
    assert body["rejected"] == [{"filename": "big.txt", "reason": "File is larger than 1 MB"}]
    assert [d["filename"] for d in body["accepted"]] == ["memo.txt"]
    assert _documents(client, inv)["memo.txt"]["status"] == "ready"


def test_mismatched_signature():
    with pytest.raises(ValidationError, match="does not match"):
        validate_file("fake.pdf", b"this is plain text, not a pdf")
    with pytest.raises(ValidationError, match="does not match"):
        validate_file("binary.txt", b"abc\x00def")


# --- chunking -----------------------------------------------------------------


def test_chunks_never_cross_pages_and_do_not_overlap():
    long_para = " ".join(f"Sentence number {i} of the long paragraph." for i in range(80))
    assert len(long_para) > MAX_PARAGRAPH_CHARS
    pages = [
        Page(1, "text", None, [Paragraph("1. First", True), Paragraph("P1A " + "a" * 500), Paragraph("P1B " + "b" * 500)]),
        Page(2, "text", None, [Paragraph("P2A " + "c" * 100), Paragraph(long_para)]),
    ]
    chunks = chunk_pages(pages)

    assert [c["ordinal"] for c in chunks] == list(range(len(chunks)))
    for c in chunks:
        on_page_1 = "P1A" in c["text"] or "P1B" in c["text"]
        on_page_2 = "P2A" in c["text"] or "Sentence number" in c["text"]
        assert not (on_page_1 and on_page_2)
        assert c["page"] == (1 if on_page_1 else 2)
        assert len(c["text"]) <= MAX_PARAGRAPH_CHARS
    # no overlap: every piece of source text appears in exactly one chunk
    for marker in ["P1A", "P1B", "P2A"] + [f"Sentence number {i} " for i in range(80)]:
        assert sum(marker in c["text"] for c in chunks) == 1
    # heading sets the section and opens the next chunk; section carries to the next page
    assert chunks[0]["text"].startswith("1. First") and chunks[0]["section"] == "1. First"
    assert chunks[0]["paragraph_index"] == 1
    assert all(c["section"] == "1. First" for c in chunks)


# --- upload + pipeline --------------------------------------------------------


def test_upload_all_formats_isolated_failures_and_duplicates(client):
    inv = _new_investigation(client)
    response = _upload(
        client,
        inv,
        [
            ("agreement.pdf", PDF),
            ("memo.txt", TXT),
            ("policy.docx", _docx_bytes()),
            ("invoice.png", _png_bytes()),
            ("agreement-copy.pdf", PDF),
            ("broken.pdf", b"%PDF-1.4\nthis is not really a pdf"),
            ("sheet.xlsx", b"PK\x03\x04data"),
        ],
    )
    assert response.status_code == 200
    body = response.json()
    assert [r["filename"] for r in body["rejected"]] == ["sheet.xlsx"]
    assert len(body["accepted"]) == 6
    assert {d["filename"]: d["status"] for d in body["accepted"]}["agreement.pdf"] == "queued"

    docs = _documents(client, inv)
    assert docs["agreement.pdf"]["status"] == "ready"
    assert docs["memo.txt"]["status"] == "ready"
    assert docs["policy.docx"]["status"] == "ready"
    assert docs["invoice.png"]["status"] == "ready"
    assert docs["agreement-copy.pdf"]["status"] == "duplicate"
    assert docs["broken.pdf"]["status"] == "failed" and docs["broken.pdf"]["error_message"]

    # PDF: page and section metadata, embeddings
    pdf = docs["agreement.pdf"]
    assert pdf["page_count"] == 2 and pdf["extraction_method"] == "text"
    chunks = _chunks(pdf["id"])
    assert len(chunks) == pdf["chunk_count"] == 2
    payment = next(c for c in chunks if "BETAPAGE" in c["text"])
    assert payment["page"] == 2 and payment["section"] == "4.2 Payment Terms"
    assert "ALPHAPAGE" not in payment["text"]
    assert next(c for c in chunks if "ALPHAPAGE" in c["text"])["page"] == 1
    for c in chunks:
        assert c["embedding"] is not None and len(c["embedding"]) == 384 * 4
        assert c["investigation_id"] == inv and len(c["text_hash"]) == 64

    # TXT and DOCX: no pages, text extracted
    txt_chunks = _chunks(docs["memo.txt"]["id"])
    assert txt_chunks[0]["page"] is None and "remote work is not permitted" in txt_chunks[0]["text"]
    assert docs["memo.txt"]["page_count"] is None
    docx_text = "\n".join(c["text"] for c in _chunks(docs["policy.docx"]["id"]))
    assert "Finance Director" in docx_text and "Late fee | 1.5% per month" in docx_text
    assert _chunks(docs["policy.docx"]["id"])[0]["section"] == "Invoice Approval"

    # image: OCR
    png = docs["invoice.png"]
    assert png["extraction_method"] == "ocr" and png["page_count"] == 1
    ocr_chunk = _chunks(png["id"])[0]
    assert ocr_chunk["extraction_method"] == "ocr" and ocr_chunk["page"] == 1
    assert 0 < ocr_chunk["ocr_confidence"] <= 1
    assert "45 days" in ocr_chunk["text"]

    # duplicate and failed documents are not indexed
    assert _chunks(docs["agreement-copy.pdf"]["id"]) == []
    assert _chunks(docs["broken.pdf"]["id"]) == []


def test_status_transitions(client, monkeypatch):
    seen = []
    original = pipeline.set_status

    def recording(document_id, status, error_message=None):
        seen.append(status)
        original(document_id, status, error_message)

    monkeypatch.setattr(pipeline, "set_status", recording)
    inv = _new_investigation(client)
    body = _upload(client, inv, [("memo.txt", TXT)]).json()
    assert body["accepted"][0]["status"] == "queued"
    assert seen == ["extracting", "indexing"]
    assert _documents(client, inv)["memo.txt"]["status"] == "ready"


def test_same_file_in_another_investigation_is_not_a_duplicate(client):
    inv = _new_investigation(client)
    _upload(client, inv, [("memo.txt", TXT)])
    assert _documents(client, inv)["memo.txt"]["status"] == "ready"


def test_upload_errors(client):
    inv = _new_investigation(client)
    assert client.post(f"/api/investigations/{inv}/documents").json()["error"]["code"] == "no_files"
    too_many = _upload(client, inv, [(f"n{i}.txt", TXT) for i in range(11)])
    assert too_many.status_code == 400 and too_many.json()["error"]["code"] == "too_many_files"
    missing = _upload(client, "does-not-exist", [("memo.txt", TXT)])
    assert missing.status_code == 404 and missing.json()["error"]["code"] == "investigation_not_found"


def test_embedder_failure_does_not_fail_the_document(client, monkeypatch):
    monkeypatch.setattr(pipeline.embedder, "embed_passages", lambda texts: None)
    inv = _new_investigation(client)
    _upload(client, inv, [("memo.txt", TXT)])
    doc = _documents(client, inv)["memo.txt"]
    assert doc["status"] == "ready"
    assert all(c["embedding"] is None for c in _chunks(doc["id"]))
