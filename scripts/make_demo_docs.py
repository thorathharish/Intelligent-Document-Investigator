"""Generate the synthetic demo documents (LLD section 19). All parties are fictional.

Usage (from repo root, venv active):
  python scripts/make_demo_docs.py --set A
"""
import argparse
from pathlib import Path

import docx
import pymupdf

ROOT = Path(__file__).resolve().parents[1]


def write_pdf(path: Path, pages: list[list[tuple[str, str]]], body_size: float = 10.5) -> None:
    """pages: list of pages; each page is a list of (kind, text) with kind in t (title), h (heading), p."""
    sizes = {"t": 16, "h": 13, "p": body_size}
    doc = pymupdf.open()
    for items in pages:
        page = doc.new_page()
        y = 72.0
        for kind, text in items:
            rect = pymupdf.Rect(72, y, 523, 800)
            unused = page.insert_textbox(
                rect, text, fontsize=sizes[kind], fontname="helv" if kind == "p" else "hebo"
            )
            y += rect.height - unused + (14 if kind == "p" else 10)
    doc.save(path)
    doc.close()


def make_set_a(out: Path) -> None:
    out.mkdir(parents=True, exist_ok=True)

    write_pdf(
        out / "Master_Services_Agreement.pdf",
        [
            [
                ("t", "MASTER SERVICES AGREEMENT"),
                ("p", "This Master Services Agreement is made between Orion Logistics Pvt. Ltd. (the Supplier) "
                      "and Kestrel Foods Ltd. (the Customer)."),
                ("h", "1. Effective Date and Term"),
                ("p", "This Agreement takes effect on 10 January 2024 and continues for a term of 24 months "
                      "unless it is terminated earlier under clause 9."),
                ("h", "2. Services"),
                ("p", "The Supplier shall provide temperature-controlled freight and warehousing services to the "
                      "Customer at the locations listed in each statement of work. The Supplier shall perform the "
                      "services with reasonable skill and care and in line with good industry practice."),
                ("h", "3. Service Levels"),
                ("p", "The Supplier shall deliver at least 97 percent of consignments within the agreed delivery "
                      "window, measured each calendar month."),
            ],
            [
                ("h", "4. Fees and Payment"),
                ("h", "4.1 Fees"),
                ("p", "The Customer shall pay the fees set out in each statement of work. Fees exclude applicable "
                      "taxes, which are charged in addition."),
                ("h", "4.2 Payment Terms"),
                ("p", "The Customer shall pay each undisputed invoice within thirty (30) days of the invoice date."),
                ("h", "4.3 Late Payment"),
                ("p", "Any amount not paid when due accrues a late fee of 1.5% per month on the outstanding balance "
                      "until it is paid in full."),
                ("h", "5. Confidentiality"),
                ("p", "Each party shall keep the other party's confidential information secret and use it only to "
                      "perform this Agreement."),
            ],
            [
                ("h", "9. Termination"),
                ("h", "9.1 Termination for Convenience"),
                ("p", "Either party may terminate this Agreement by giving 60 days' written notice to the other party."),
                ("h", "9.2 Termination for Breach"),
                ("p", "Either party may terminate this Agreement immediately if the other party commits a material "
                      "breach and fails to remedy it after being asked in writing to do so."),
                ("h", "10. Governing Law"),
                ("p", "This Agreement is governed by the laws of India."),
            ],
        ],
    )

    write_pdf(
        out / "Amendment_1.pdf",
        [
            [
                ("t", "AMENDMENT NO. 1 TO MASTER SERVICES AGREEMENT"),
                ("p", "This Amendment No. 1 is dated 15 February 2024 and amends the Master Services Agreement "
                      "between Orion Logistics Pvt. Ltd. and Kestrel Foods Ltd. that took effect on 10 January 2024 "
                      "(the Agreement)."),
                ("h", "1. Payment Terms"),
                ("p", "Clause 4.2 of the Agreement remains unchanged. The Customer shall pay each undisputed "
                      "invoice within 30 days of the invoice date."),
                ("h", "2. Termination Notice"),
                ("p", "This Amendment replaces clause 9.1 of the Agreement. Either party may terminate the "
                      "Agreement by giving 90 days' written notice to the other party."),
                ("h", "3. Other Terms"),
                ("p", "All other terms of the Agreement remain in full force and effect."),
            ]
        ],
    )

    # The invoice is produced twice: a text PDF (used until OCR is in place) and a PNG "scan" of it.
    invoice_pdf = out / "Invoice_INV-2041.pdf"
    write_pdf(
        invoice_pdf,
        [
            [
                ("t", "INVOICE INV-2041"),
                ("p", "From Orion Logistics Pvt. Ltd. to Kestrel Foods Ltd."),
                ("p", "This invoice is dated 1 March 2024."),
                ("p", "Cold-chain freight services for February 2024."),
                ("p", "The total amount due is USD 48,500."),
                ("p", "Payment due within 45 days of the invoice date."),
            ]
        ],
        body_size=13,
    )
    with pymupdf.open(invoice_pdf) as doc:
        doc[0].get_pixmap(dpi=150).save(out / "Invoice_INV-2041.png")

    policy = docx.Document()
    policy.add_heading("Vendor Payment Policy", level=1)
    policy.add_paragraph("This policy applies to all vendor invoices received by Kestrel Foods Ltd.")
    policy.add_heading("Invoice Approval", level=2)
    policy.add_paragraph("Invoices above USD 25,000 must be approved by the Finance Director before payment.")
    policy.add_heading("Late Fees", level=2)
    policy.add_paragraph(
        "Kestrel Foods accepts a late fee of 1.5% per month on overdue vendor invoices, in line with its "
        "standard supplier contracts."
    )
    policy.add_heading("Records", level=2)
    policy.add_paragraph("Approved invoices and proof of payment are kept for seven years.")
    policy.save(str(out / "Vendor_Payment_Policy.docx"))

    print("Set A written to", out)
    for f in sorted(out.iterdir()):
        print("  ", f.name)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--set", choices=["A"], required=True)
    args = parser.parse_args()
    if args.set == "A":
        make_set_a(ROOT / "demo_docs" / "set_a")


if __name__ == "__main__":
    main()
